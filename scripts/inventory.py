"""Walk materials/ and write bank/inventory.json and docs/data/datasets.json.

Never executes macros: xlsm files are read as data only and the VBA project is
only inspected as a zip member.
"""
import csv
import hashlib
import io
import json
import re
import sqlite3
import zipfile
from collections import Counter
from pathlib import Path

import docx
import openpyxl
import pptx

ROOT = Path(__file__).resolve().parent.parent
MAT = ROOT / "materials"
SAMPLE_ROWS = 200  # rows sampled for column types

# Classification by file name. Everything else falls back to the extension rule.
CLASS_BY_NAME = {
    "solutions.xlsx": "solution",
    "PWC answers.docx": "solution",
    "PWC.pptx": "slides",
    "Note.txt": "note",
    "requierements.txt": "instructions",
    "requirements.txt": "instructions",
    "Instructions.docx": "instructions",
    "Assignment_01_Excel-SQL.xlsx": "instructions",
}
CLASS_BY_EXT = {".docx": "instructions", ".pptx": "slides", ".xlsx": "dataset", ".xlsm": "dataset",
                ".csv": "dataset", ".txt": "dataset", ".db": "dataset", ".sqlite": "dataset"}

CASE_BY_DIR = {
    "1. Excel Part 1": "Excel Part 1",
    "1. Excel part 2": "Excel Part 2",
    "1. Excel part 3": "Excel Part 3",
    "1. Excel part 4": "Excel Part 4",
    "2. SQL part 2": "SQL Part 2 P-card",
}


def case_of(rel):
    parts = rel.split("/")
    if parts[0] == "2. SQL part 1":
        return "SQL Part 1 PWC" if "pwc assignment" in parts else "SQL Part 1 CS50"
    return CASE_BY_DIR.get(parts[0], parts[0])


def sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


# ---------- docx / pptx ----------

def read_docx(path):
    d = docx.Document(path)
    lines = []
    for p in d.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        style = p.style.name if p.style is not None else ""
        numbered = p._p.pPr is not None and p._p.pPr.numPr is not None
        if style.startswith("Heading"):
            lines.append("#" * int(style[-1] if style[-1].isdigit() else 1) + " " + text)
        elif numbered or "List" in style:
            level = 0
            if numbered and p._p.pPr.numPr.ilvl is not None:
                level = int(p._p.pPr.numPr.ilvl.val)
            lines.append("  " * level + "- " + text)
        else:
            lines.append(text)
    tables = [[[c.text.strip() for c in row.cells] for row in t.rows] for t in d.tables]
    images = sum(1 for r in d.part.rels.values() if "image" in r.reltype)
    return {"text": "\n".join(lines), "tables": tables, "images": images}


def read_pptx(path):
    prs = pptx.Presentation(path)
    slides = []
    for i, s in enumerate(prs.slides, 1):
        texts = []
        for sh in s.shapes:
            if sh.has_text_frame:
                t = sh.text_frame.text.strip()
                if t:
                    texts.append(t)
            if getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    texts.append(" | ".join(c.text for c in row.cells))
        notes = s.notes_slide.notes_text_frame.text.strip() if s.has_notes_slide else ""
        slides.append({"slide": i, "text": "\n".join(texts), "notes": notes})
    return {"slides": slides}


# ---------- spreadsheets ----------

def cell_type(v):
    if v is None:
        return None
    if isinstance(v, bool):
        return "bool"
    if isinstance(v, (int, float)):
        return "number"
    if hasattr(v, "year"):
        return "date"
    if isinstance(v, str) and v.startswith("="):
        return "formula"
    return "text"


def detect_header(rows):
    """First all-text row (within the first 10) holding at least 60 % of the widest such row's cells.

    Skips title rows above the real header, e.g. a merged 'Solution using formulas' label or a description line.
    """
    if not rows:
        return None
    width = lambda r: sum(1 for v in r if v not in (None, ""))
    text_rows = [i for i, r in enumerate(rows[:10])
                 if width(r) >= 2 and all(isinstance(v, str) for v in r if v not in (None, ""))]
    if not text_rows:
        return 0
    widest = max(width(rows[i]) for i in text_rows)
    return next(i for i in text_rows if width(rows[i]) >= 0.6 * widest)


def column_types(rows, header_idx, width):
    types = []
    data = rows[header_idx + 1: header_idx + 1 + SAMPLE_ROWS]
    for c in range(width):
        cnt = Counter(cell_type(r[c]) for r in data if c < len(r))
        cnt.pop(None, None)
        types.append("/".join(t for t, _ in cnt.most_common()) or "empty")
    return types


def read_workbook(path):
    wb = openpyxl.load_workbook(path, data_only=False)
    sheets = []
    for ws in wb.worksheets:
        rows = [tuple(r) for r in ws.iter_rows(values_only=True)]
        while rows and all(v is None for v in rows[-1]):
            rows.pop()
        formulas = sum(1 for r in rows for v in r if isinstance(v, str) and v.startswith("="))
        h = detect_header(rows)
        header = [v for v in rows[h]] if h is not None else []
        while header and header[-1] is None:
            header.pop()
        sheets.append({
            "name": ws.title,
            "state": ws.sheet_state,
            "dimensions": ws.dimensions,
            "max_row": ws.max_row,
            "max_column": ws.max_column,
            "header_row": None if h is None else h + 1,
            "header": [None if v is None else str(v) for v in header],
            "data_rows": 0 if h is None else sum(1 for r in rows[h + 1:] if any(v is not None for v in r)),
            "column_types": column_types(rows, h, len(header)) if h is not None else [],
            "filled_columns": [] if h is None else [
                str(header[c]) for c in range(len(header))
                if header[c] is not None and any(c < len(r) and r[c] not in (None, "") for r in rows[h + 1:])],
            "formula_cells": formulas,
            "formula_examples": [v for r in rows for v in r if isinstance(v, str) and v.startswith("=")][:5],
            "merged_cells": [str(m) for m in ws.merged_cells.ranges][:50],
            "first_rows": [[None if v is None else str(v)[:60] for v in r[:15]] for r in rows[:8]],
        })
    out = {"sheets": sheets, "defined_names": list(wb.defined_names.keys())}
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        vba = [n for n in names if n.endswith("vbaProject.bin")]
        out["has_vba_project"] = bool(vba)
        if vba:
            out["vba_modules"] = vba_module_names(z.read(vba[0]))
    return out


def vba_module_names(blob):
    """Module names from the OLE directory of vbaProject.bin, without running anything."""
    try:
        import olefile
    except ImportError:
        # Fallback: stream names are stored as UTF-16 in the OLE directory entries.
        names = set(re.findall(rb"((?:[A-Za-z0-9_]\x00){3,31})\x00\x00", blob))
        names = {n.decode("utf-16-le") for n in names}
        skip = {"Root Entry", "VBA", "PROJECT", "PROJECTwm", "_VBA_PROJECT", "dir"}
        return sorted(n for n in names if n not in skip and not n.startswith("__SRP"))[:30]
    ole = olefile.OleFileIO(io.BytesIO(blob))
    return sorted("/".join(s) for s in ole.listdir() if s[0] == "VBA" and s[-1] not in ("dir", "_VBA_PROJECT"))


# ---------- csv / txt ----------

def read_text_file(path):
    raw = path.read_bytes()
    bom = raw.startswith(b"\xef\xbb\xbf")
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    encoding = "utf-8" if enc == "utf-8-sig" else enc
    lines = text.splitlines()
    out = {"encoding": encoding, "bom": bom, "line_endings": "CRLF" if b"\r\n" in raw[:100000] else "LF",
           "lines": len(lines)}
    # A delimiter counts if most of the first lines split into the same number (>1) of fields.
    head = [l for l in lines[:50] if l.strip()]
    best = None
    for d in [",", "|", ";", "\t"]:
        counts = Counter(len(next(csv.reader([l], delimiter=d))) for l in head)
        if not counts:
            continue
        width, hits = counts.most_common(1)[0]
        if width > 1 and hits >= 0.8 * len(head) and (best is None or (hits, width) > best[2:]):
            best = (d, width, hits, width)
    if best is None:
        out["kind"] = "text"
        out["text"] = text if len(text) < 20000 else text[:20000]
        return out
    d = best[0]
    rows = list(csv.reader(io.StringIO(text), delimiter=d))
    rows = [r for r in rows if any(c.strip() for c in r)]
    has_header = not any(re.fullmatch(r"-?[\d.,$]+", c.strip()) for c in rows[0] if c.strip())
    widths = Counter(len(r) for r in rows)
    out.update({
        "kind": "delimited",
        "delimiter": {",": "comma", "|": "pipe", ";": "semicolon", "\t": "tab"}[d],
        "has_header": has_header,
        "header": rows[0] if has_header else [],
        "columns": best[1],
        "row_widths": dict(widths),
        "data_rows": len(rows) - (1 if has_header else 0),
        "sample": rows[1:6] if has_header else rows[:5],
    })
    return out


# ---------- sqlite ----------

def read_sqlite(path):
    con = sqlite3.connect(f"file:{path.as_posix()}?mode=ro", uri=True)
    objs = []
    for name, typ in con.execute("SELECT name, type FROM sqlite_master WHERE type IN ('table','view') "
                                 "AND name NOT LIKE 'sqlite_%' ORDER BY type, name"):
        cols = [{"name": c[1], "type": c[2], "pk": bool(c[5])}
                for c in con.execute(f'PRAGMA table_info("{name}")')]
        fks = [{"column": f[3], "references": f"{f[2]}.{f[4]}"}
               for f in con.execute(f'PRAGMA foreign_key_list("{name}")')]
        n = con.execute(f'SELECT COUNT(*) FROM "{name}"').fetchone()[0]
        objs.append({"name": name, "type": typ, "columns": cols, "foreign_keys": fks, "rows": n})
    con.close()
    return {"objects": objs}


# ---------- main ----------

READERS = {".docx": read_docx, ".pptx": read_pptx, ".xlsx": read_workbook, ".xlsm": read_workbook,
           ".csv": read_text_file, ".txt": read_text_file, ".db": read_sqlite, ".sqlite": read_sqlite}


def classify(path, content):
    cls = CLASS_BY_NAME.get(path.name, CLASS_BY_EXT.get(path.suffix.lower(), "note"))
    notes = []
    if path.suffix.lower() == ".txt" and content.get("kind") == "text" and cls == "dataset":
        cls = "note"
    # Solution content mixed into a non-solution file: worksheets or sheets named like answers.
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        for s in content.get("sheets", []):
            if re.search(r"solution|answer|output|result|regression|summary", s["name"], re.I):
                filled = s["filled_columns"]
                notes.append(f"sheet '{s['name']}' is named like a solution: {len(filled)} of {len(s['header'])} "
                             f"header columns hold data ({', '.join(filled)})")
            if s["formula_cells"]:
                notes.append(f"sheet '{s['name']}' has {s['formula_cells']} formula cells")
            answer_cols = [h for h in s["header"] if h and re.fullmatch(r"answers?", h, re.I)]
            if answer_cols:
                notes.append(f"sheet '{s['name']}' has an Answer column; filled: "
                             f"{'Answer' in s['filled_columns'] or 'Answers' in s['filled_columns']}")
    return cls, notes


def main():
    entries = []
    for p in sorted(MAT.rglob("*")):
        if not p.is_file():
            continue
        rel = p.relative_to(MAT).as_posix()
        print("reading", rel, flush=True)
        e = {"path": rel, "case": case_of(rel), "ext": p.suffix.lower(), "size_bytes": p.stat().st_size,
             "sha256": sha256(p)}
        try:
            e["content"] = READERS[p.suffix.lower()](p)
        except Exception as ex:  # record and keep going; the notes list unreadable files
            e["content"] = {}
            e["error"] = f"{type(ex).__name__}: {ex}"
        e["class"], e["class_notes"] = classify(p, e["content"])
        entries.append(e)
    overrides = ROOT / "bank" / "class_overrides.json"
    if overrides.exists():
        for e in entries:
            o = json.loads(overrides.read_text(encoding="utf-8")).get(e["path"])
            if o:
                e["class"], e["class_reason"] = o["class"], o["reason"]
    (ROOT / "bank" / "inventory.json").write_text(json.dumps(entries, indent=1, ensure_ascii=False, default=str),
                                                  encoding="utf-8")
    write_datasets(entries)
    print(f"{len(entries)} files inventoried")


DATASET_LABELS = json.loads((ROOT / "scripts" / "dataset_labels.json").read_text(encoding="utf-8"))


def write_datasets(entries):
    out = []
    for e in entries:
        if e["class"] != "dataset" or e["path"] not in DATASET_LABELS:
            continue
        meta = DATASET_LABELS[e["path"]]
        d = {"id": meta["id"], "label": meta["label"], "case": e["case"], "file": Path(e["path"]).name,
             "size_bytes": e["size_bytes"]}
        c = e["content"]
        if "sheets" in c:
            d["sheets"] = [{"name": s["name"], "header_row": s["header_row"], "headers": s["header"],
                            "rows": s["data_rows"]} for s in c["sheets"] if s["state"] == "visible"]
        elif "objects" in c:
            d["tables"] = [{"name": o["name"], "type": o["type"], "rows": o["rows"],
                            "columns": [f"{col['name']} {col['type']}".strip() for col in o["columns"]]}
                           for o in c["objects"]]
        else:
            d["delimiter"] = c.get("delimiter")
            d["encoding"] = c.get("encoding")
            d["headers"] = c.get("header")
            d["rows"] = c.get("data_rows")
        out.append(d)
    (ROOT / "docs" / "data" / "datasets.json").write_text(json.dumps(out, indent=1, ensure_ascii=False),
                                                          encoding="utf-8")


if __name__ == "__main__":
    main()
